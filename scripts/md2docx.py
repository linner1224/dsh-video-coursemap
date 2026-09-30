#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""md2docx.py <in.md> <out.docx> — 把受控 Markdown 转成 Word 文档（.docx）。

支持：标题 #/##/###、无序列表 -（含缩进层级）、引用 >、```代码块```、
**加粗** / *斜体* 内联；忽略 ---、HTML 注释、空行。

特别地：```mermaid ... mindmap ... ``` 代码块会被解析成原生 Word 分级大纲
（根节点加粗 + 子节点分级项目符号），使「知识地图」在 Word 里呈现为真正的
层级结构，而非原始代码文本。
"""
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

_TOKEN_RE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*)')


def _set_east_asia(run, east='宋体'):
    """给 run 设置东亚字体，避免中文在个别 Word 版本里显示异常。"""
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn('w:eastAsia'), east)


def add_inline(paragraph, text):
    """解析 **bold** / *italic* 并逐段添加 run。"""
    for tok in _TOKEN_RE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) >= 4:
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) >= 2 and not tok.startswith('**'):
            r = paragraph.add_run(tok[1:-1])
            r.italic = True
        else:
            paragraph.add_run(tok)


def _mindmap_text(s):
    """清洗 mindmap 节点行：去掉 root 前缀与 (( ))/[ ]/( )/{ } 包裹，返回纯文字。"""
    t = s.strip()
    t = re.sub(r'^root\s*', '', t)
    for _ in range(3):
        if t.startswith('((') and t.endswith('))'):
            t = t[2:-2].strip()
        elif t.startswith('))') and t.endswith('(('):
            t = t[2:-2].strip()
        elif t.startswith('[') and t.endswith(']'):
            t = t[1:-1].strip()
        elif t.startswith('(') and t.endswith(')'):
            t = t[1:-1].strip()
        elif t.startswith('{') and t.endswith('}'):
            t = t[1:-1].strip()
        else:
            break
    return t


def _is_mindmap(code_lines):
    first = next((ln.strip() for ln in code_lines if ln.strip()), '')
    return first.lower().startswith('mindmap')


def _render_mindmap(doc, lines):
    """把 mindmap 代码块渲染为 Word 分级大纲（根节点加粗，子节点分级列表）。"""
    nodes = []
    for ln in lines:
        s = ln.rstrip()
        if not s.strip():
            continue
        if s.strip().lower() == 'mindmap':
            continue
        indent = len(s) - len(s.lstrip(' '))
        text = _mindmap_text(s)
        if not text:
            continue
        nodes.append((indent, text))
    if not nodes:
        return
    # 根节点：最小缩进 → 加粗段落
    root_indent = min(i for i, _ in nodes)
    root_idx = next(k for k, (i, _) in enumerate(nodes) if i == root_indent)
    p = doc.add_paragraph()
    r = p.add_run(nodes[root_idx][1])
    r.bold = True
    # 其余节点：相对缩进 → 分级项目符号（最多 3 级）
    rest = [(i, t) for k, (i, t) in enumerate(nodes) if k != root_idx]
    if not rest:
        return
    indents = sorted(set(i for i, _ in rest))
    ind2lvl = {i: min(idx, 2) for idx, i in enumerate(indents)}
    styles = ['List Bullet', 'List Bullet 2', 'List Bullet 3']
    for i, t in rest:
        add_inline(doc.add_paragraph(style=styles[ind2lvl[i]]), t)


def convert(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()
    # Normal 样式基础字体（拉丁 + 东亚）
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')

    i = 0
    in_code = False
    code_lines = []

    def flush_code():
        nonlocal in_code, code_lines
        if code_lines:
            if _is_mindmap(code_lines):
                _render_mindmap(doc, code_lines)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(18)
                for j, cl in enumerate(code_lines):
                    run = p.add_run((cl if j == 0 else '\n' + cl))
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
        in_code = False
        code_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith('```'):
            if in_code:
                flush_code()
            else:
                in_code = True
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if stripped == '' or stripped == '---':
            i += 1
            continue
        if stripped.startswith('<!--') and stripped.endswith('-->'):
            i += 1
            continue

        if stripped.startswith('### '):
            h = doc.add_heading('', level=3)
            add_inline(h, stripped[4:])
        elif stripped.startswith('## '):
            h = doc.add_heading('', level=2)
            add_inline(h, stripped[3:])
        elif stripped.startswith('# '):
            h = doc.add_heading('', level=1)
            add_inline(h, stripped[2:])
        elif stripped.startswith('- '):
            indent = len(line) - len(line.lstrip(' '))
            level = min(indent // 2 + 1, 3)
            style = 'List Bullet' if level == 1 else ('List Bullet 2' if level == 2 else 'List Bullet 3')
            p = doc.add_paragraph(style=style)
            add_inline(p, stripped[2:])
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            r = p.add_run(stripped[2:])
            r.italic = True
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        else:
            p = doc.add_paragraph()
            add_inline(p, stripped)
        i += 1

    if in_code:
        flush_code()

    doc.save(docx_path)
    print('OK', docx_path)


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print('usage: md2docx.py <in.md> <out.docx>', file=sys.stderr)
        sys.exit(2)
    convert(sys.argv[1], sys.argv[2])
